import openpyxl
from datetime import datetime

today = datetime.now().strftime("%Y-%m-%d")

# 1. Update Automation Test Cases Workbook
auto_path = "docs/Verity_Automation_Test_Cases.xlsx"
wb_auto = openpyxl.load_workbook(auto_path)
ws_prof_auto = wb_auto["Profile_Personal"]

for row in ws_prof_auto.iter_rows(min_row=2):
    test_id = row[0].value
    if test_id == "AUT-PROF-004":
        row[3].value = "Tab-Level Reset and Save Controls"
        row[4].value = "Verify each profile tab has working Reset and Save controls, Reset clears all editable fields and drafts to empty state without backend mutation, and Save persists changes cleanly"
        row[9].value = "1. Fill fields in Personal Info. 2. Verify active badge. 3. Click Reset -> Cancel preserves fields. 4. Click Reset -> Confirm Clear Fields -> all fields become empty strings. 5. Save persists valid edits and updates baseline."
        row[10].value = "Reset clears all input fields, textareas, dropdowns, and edit drafts to empty; does not delete saved DB records or Master Resume PDF; Save validates and persists"
        row[11].value = "await expect(headlineInput).toHaveValue('') and await expect(page.getByText('Cleared all editable fields in Personal Info')).toBeVisible()"
        row[13].value = "Passed"
        row[16].value = today
        row[17].value = "Verified passing: Reset clears fields to empty state without deleting DB data"
    elif test_id == "AUT-PROF-005":
        row[3].value = "Global Reset All Profile Controls"
        row[4].value = "Verify Reset All button prompts confirmation when active forms exist, clears editable input fields across all 6 tabs to empty state, and preserves database records and Master Resume files"
        row[9].value = "1. Enter draft inputs in multiple tabs. 2. Click Reset All. 3. Verify confirmation modal appears. 4. Confirm Clear All Forms. 5. Verify all forms cleared to empty defaults and active tab preserved."
        row[10].value = "Reset All clears all editable form fields and drafts across all tabs; no backend delete or reset requests sent; database records and Master Resume files preserved"
        row[11].value = "await expect(skillInput).toHaveValue('') and await expect(page.getByText('Cleared all editable fields across all Profile tabs')).toBeVisible()"
        row[13].value = "Passed"
        row[16].value = today
        row[17].value = "Verified passing: Reset All clears all forms across all tabs"

wb_auto.save(auto_path)
print("Automation workbook updated.")

# 2. Update Manual Test Cases Workbook
manual_path = "docs/Verity_Manual_Test_Cases.xlsx"
wb_manual = openpyxl.load_workbook(manual_path)
ws_prof_manual = wb_manual["Profile_Personal"]

for row in ws_prof_manual.iter_rows(min_row=2):
    test_id = row[0].value
    if test_id == "TC-PROF-004":
        row[2].value = "Profile_Personal"
        row[3].value = "Verify candidate can reset and clear all editable input fields and drafts on any tab, leaving them empty without deleting saved database records"
        row[6].value = "1. Open any editable profile tab.\n2. Enter text into input fields or edit a card.\n3. Click Reset -> verify confirmation modal appears.\n4. Click Cancel -> verify fields remain filled.\n5. Click Clear Fields -> verify all editable input fields are cleared and empty.\n6. Verify saved database cards and Master Resume PDF are intact.\n7. Enter valid details and click Save -> verify save succeeds."
        row[7].value = "Reset clears all inputs, dropdowns, dates, and drafts to empty defaults without sending backend delete requests; Save validates and persists normally."
        row[12].value = "Reset clears all inputs to empty strings and cancels active edit sessions; preserved saved records and Master Resume files."
        row[13].value = "Passed"
        row[16].value = today
        row[18].value = "Verified tab Reset clears form fields to empty state"
    elif test_id == "TC-PROF-005":
        row[2].value = "Profile_Personal"
        row[3].value = "Verify candidate can clear all editable input fields and drafts across all Profile tabs simultaneously using Reset All"
        row[6].value = "1. Make edits or enter draft details in multiple tabs.\n2. Click 'Reset All' button near Candidate Profile heading.\n3. Verify confirmation modal explains input fields will be cleared and saved records will not be deleted.\n4. Click Cancel -> verify all fields remain filled.\n5. Click Clear All Forms -> verify all forms across all 6 tabs are cleared to empty state.\n6. Verify saved database records and Master Resume files remain intact."
        row[7].value = "Reset All clears all form fields across all tabs to empty defaults after confirmation without deleting saved profile data or Master Resume PDF."
        row[12].value = "Reset All clears all forms to empty state; preserves DB records and Master Resume PDF."
        row[13].value = "Passed"
        row[16].value = today
        row[18].value = "Verified Reset All clears all forms to empty state"

wb_manual.save(manual_path)
print("Manual workbook updated.")
