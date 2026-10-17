import openpyxl
from datetime import datetime

today = datetime.now().strftime("%Y-%m-%d")

# 1. Update Automation Test Cases
auto_path = "docs/PrivateResume_Automation_Test_Cases.xlsx"
wb_auto = openpyxl.load_workbook(auto_path)
ws_jda_auto = wb_auto["JD_Analysis"]

jda_tests = [
    (
        "AUT-JDA-001", "TC-JDA-001", "JD_Analysis", "Ollama Abstraction", "Deterministic & local requirement extraction from JD text",
        "Unit", "Pytest / Provider", "Sample JD text with FastAPI, React, PostgreSQL", "Raw JD string",
        "Call deterministic_jd_extraction", "Extracts languages, frameworks, tools, cloud, and experience level",
        "assert 'FastAPI' in analysis.frameworks and '5+ Years' in analysis.experience_years",
        "backend/tests/test_jd_analysis.py::test_ollama_provider_deterministic_extraction",
        "Passed", "", "Local / Epic-4", today, "Verified locally"
    ),
    (
        "AUT-JDA-002", "TC-JDA-002", "JD_Analysis", "Trigger Analysis", "Trigger JD analysis and persist JobRequirement record",
        "API", "Pytest / TestClient", "Job created in DB", "POST /api/v1/jobs/{id}/analyze",
        "Execute analysis endpoint and verify persisted response", "Returns 200 OK and structured JobAnalysisResponse",
        "assert response.status_code == 200 and 'Python' in data['languages']",
        "backend/tests/test_jd_analysis.py::test_analyze_job_and_retrieve_analysis",
        "Passed", "", "Local / Epic-4", today, "Verified locally"
    ),
    (
        "AUT-JDA-003", "TC-JDA-003", "JD_Analysis", "Validation / Negative", "Reject analysis request for nonexistent job",
        "API / Negative", "Pytest / TestClient", "None", "job_id=99999",
        "POST /api/v1/jobs/99999/analyze", "Returns 404 Not Found",
        "assert response.status_code == 404",
        "backend/tests/test_jd_analysis.py::test_analyze_nonexistent_job_fails",
        "Passed", "", "Local / Epic-4", today, "Verified locally"
    ),
    (
        "AUT-JDA-004", "TC-JDA-004", "JD_Analysis", "Security / Isolation", "Cross-user job analysis access rejection",
        "Security", "Pytest / TestClient", "Two distinct registered candidates", "User 1 and User 2 tokens",
        "User 2 attempts GET and POST analyze on User 1 job", "Returns 404 Not Found across both operations",
        "assert 404 on cross-user analysis access attempts",
        "backend/tests/test_jd_analysis.py::test_job_analysis_user_isolation",
        "Passed", "", "Local / Epic-4", today, "Verified locally"
    ),
]

for row in jda_tests:
    ws_jda_auto.append(row)

# Update Index in Auto workbook
idx_ws_auto = wb_auto["00_Index"]
for row in idx_ws_auto.iter_rows(min_row=2):
    if row[0].value == "JD_Analysis":
        row[2].value = len(jda_tests)
        row[3].value = today
        row[4].value = len(jda_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "4 automated tests passed (Unit, API, Security)"
        break

wb_auto.save(auto_path)

# 2. Update Manual Test Cases
man_path = "docs/PrivateResume_Manual_Test_Cases.xlsx"
wb_man = openpyxl.load_workbook(man_path)
ws_jda_man = wb_man["JD_Analysis"]

jda_man_tests = [
    (
        "TC-JDA-001", "JD_Analysis", "Analysis Display", "View structured job requirements in Stage 2",
        "Target job saved in Stage 1", "Raw JD submitted",
        "1. Submit job in Stage 1. 2. Observe Stage 2 requirements breakdown.",
        "Categorized chips for Required/Preferred skills, languages, frameworks, responsibilities displayed",
        "High", "P1", "Functional", "Local Dev",
        "Stage 2 rendered all technical categories and responsibilities cleanly", "Passed", "", "QA Engineer", today, "", "Verified in browser"
    ),
    (
        "TC-JDA-002", "JD_Analysis", "Re-analysis", "Trigger re-analysis from UI",
        "Stage 2 active with existing analysis", "None",
        "1. Click 'Re-analyze JD' button.",
        "Analysis re-triggered through Ollama provider and view refreshed",
        "Medium", "P2", "Functional", "Local Dev",
        "Re-analysis completed and success banner confirmed update", "Passed", "", "QA Engineer", today, "", "Verified in browser"
    ),
]

for row in jda_man_tests:
    ws_jda_man.append(row)

# Update Index in Manual workbook
idx_ws_man = wb_man["00_Index"]
for row in idx_ws_man.iter_rows(min_row=2):
    if row[0].value == "JD_Analysis":
        row[2].value = len(jda_man_tests)
        row[3].value = today
        row[4].value = len(jda_man_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "2 manual test cases executed and passed"
        break

wb_man.save(man_path)
print("Updated Epic 4 test cases successfully!")
