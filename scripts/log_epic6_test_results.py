import openpyxl
from datetime import datetime

today = datetime.now().strftime("%Y-%m-%d")

# 1. Update Automation Test Cases Workbook
auto_path = "docs/PrivateResume_Automation_Test_Cases.xlsx"
wb_auto = openpyxl.load_workbook(auto_path)

# Sheet: Tailored_Resume
ws_tailor_auto = wb_auto["Tailored_Resume"]
tailor_auto_tests = [
    (
        "AUT-TAILOR-001", "TC-TAILOR-001", "Tailored_Resume", "Evidence Grounding", "Draft initial tailored resume from candidate evidence",
        "API / Integration", "Pytest / TestClient", "Candidate profile with Python, FastAPI, Docker; Job with Python, FastAPI, Kubernetes",
        "POST /api/v1/jobs/{id}/tailor", "Generates draft v1 resume populated with real candidate evidence",
        "assert res.status_code == 201 and data['version_number'] == 1 and 'summary' in data['content']",
        "backend/tests/test_tailoring.py::test_tailor_resume_and_non_fabrication",
        "Passed", "", "Local / Epic-6", today, "Verified locally"
    ),
    (
        "AUT-TAILOR-002", "TC-TAILOR-002", "Tailored_Resume", "Non-Fabrication Guardrail", "Ensure missing JD skills are excluded from resume",
        "Unit / API", "Pytest / TestClient", "Job JD requires Kubernetes; candidate does NOT possess Kubernetes",
        "POST /api/v1/jobs/{id}/tailor", "Resume contains Python & FastAPI; strictly does NOT contain Kubernetes",
        "assert 'Kubernetes' not in all_skills",
        "backend/tests/test_tailoring.py::test_tailor_resume_and_non_fabrication",
        "Passed", "", "Local / Epic-6", today, "Verified locally"
    ),
    (
        "AUT-TAILOR-003", "TC-TAILOR-003", "Tailored_Resume", "User Isolation", "Block cross-user tailoring and version access",
        "Security", "Pytest / TestClient", "Two distinct users with saved jobs",
        "User 2 attempts GET/POST on User 1 tailored resume", "Returns 404 Not Found",
        "assert client.get('/tailored-resumes/{id}', headers=u2_headers).status_code == 404",
        "backend/tests/test_tailoring.py::test_tailored_resume_security_isolation",
        "Passed", "", "Local / Epic-6", today, "Verified locally"
    ),
]
for row in tailor_auto_tests:
    ws_tailor_auto.append(row)

# Sheet: AI_Editing
ws_ai_auto = wb_auto["AI_Editing"]
ai_auto_tests = [
    (
        "AUT-AIEDIT-001", "TC-AIEDIT-001", "AI_Editing", "Section Refinement", "Request localized AI rewrite with anti-fabrication enforcement",
        "API / Service", "Pytest / TestClient", "Tailored draft resume v1",
        "POST /api/v1/tailored-resumes/{id}/edit with section='experience'", "Returns suggested_text, explanation, and factuality_checked=True",
        "assert data['factuality_checked'] is True and len(data['suggested_text']) > 10",
        "backend/tests/test_tailoring.py::test_ask_ai_section_editing",
        "Passed", "", "Local / Epic-6", today, "Verified locally"
    ),
]
for row in ai_auto_tests:
    ws_ai_auto.append(row)

# Sheet: Versioning
ws_ver_auto = wb_auto["Versioning"]
ver_auto_tests = [
    (
        "AUT-VER-001", "TC-VER-001", "Versioning", "Historical Immutability", "Branch new version when editing or saving snapshot",
        "Integration", "Pytest / TestClient", "Existing draft v1 resume",
        "PUT /api/v1/tailored-resumes/{id} with save_as_new_version=True", "Creates v2 resume; v1 remains unchanged",
        "assert v2['version_number'] == 2 and v2['parent_version_id'] == v1['id']",
        "backend/tests/test_tailoring.py::test_tailored_resume_versioning_and_immutability",
        "Passed", "", "Local / Epic-6", today, "Verified locally"
    ),
    (
        "AUT-VER-002", "TC-VER-002", "Versioning", "Final Lock Immutability", "Editing finalized resume forces automatic new version",
        "Integration", "Pytest / TestClient", "Finalized v2 resume",
        "PUT /api/v1/tailored-resumes/{id} on finalized resume", "Forces creation of draft v3 instead of mutating locked v2",
        "assert v3['version_number'] == 3 and v3['status'] == 'draft'",
        "backend/tests/test_tailoring.py::test_tailored_resume_versioning_and_immutability",
        "Passed", "", "Local / Epic-6", today, "Verified locally"
    ),
]
for row in ver_auto_tests:
    ws_ver_auto.append(row)

# Update Auto Index counts
idx_auto = wb_auto["00_Index"]
for row in idx_auto.iter_rows(min_row=2):
    mod = row[0].value
    if mod == "Tailored_Resume":
        row[2].value = 3
        row[3].value = today
        row[4].value = 3
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (3 automated tests passed)"
    elif mod == "AI_Editing":
        row[2].value = 1
        row[3].value = today
        row[4].value = 1
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (1 automated test passed)"
    elif mod == "Versioning":
        row[2].value = 2
        row[3].value = today
        row[4].value = 2
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (2 automated tests passed)"

wb_auto.save(auto_path)
print("Updated Automation Test Cases Workbook.")

# 2. Update Manual Test Cases Workbook
man_path = "docs/PrivateResume_Manual_Test_Cases.xlsx"
wb_man = openpyxl.load_workbook(man_path)

# Sheet: Tailored_Resume
ws_tailor_man = wb_man["Tailored_Resume"]
tailor_man_tests = [
    (
        "TC-TAILOR-001", "Tailored_Resume", "Tailored Resume Generation", "Generate v1 resume draft for target job from workspace",
        "Job added and analyzed in Stages 1 & 2", "Navigate to Stage 4 and click 'Generate Tailored Resume (Draft v1)'",
        "Target Job: Lead Developer", "Draft v1 is generated with summary, skills, experience, and projects populated",
        "Draft v1 displayed in Workspace Stage 4 with all sections editable", "Passed", "", "Manual / UI", "High",
        "Lead Engineer", today, "Verified in browser"
    ),
    (
        "TC-TAILOR-002", "Tailored_Resume", "Non-Fabrication Rule", "Verify missing skills are excluded from generated resume",
        "Job description contains skills candidate does not have", "Inspect Skills section in tailored draft",
        "Candidate skills vs Job skills", "Missing skills identified in Stage 3 are NOT added to the resume",
        "Missing skills absent from generated resume; only genuine skills present", "Passed", "", "Manual / UI", "Critical",
        "Lead Engineer", today, "Verified in browser"
    ),
]
for row in tailor_man_tests:
    ws_tailor_man.append(row)

# Sheet: AI_Editing
ws_ai_man = wb_man["AI_Editing"]
ai_man_tests = [
    (
        "TC-AIEDIT-001", "AI_Editing", "Ask AI Section Revision", "Refine summary and experience bullets with AI consultation modal",
        "Tailored resume draft loaded in Stage 4", "Click 'Ask AI' button, select prompt pill, and click Generate",
        "Prompt: 'Emphasize measurable outcomes and technical metrics'", "AI returns proposed revision with clear rationale",
        "Proposed revision displayed side-by-side; clicking Accept updates the resume section", "Passed", "", "Manual / UI", "High",
        "Lead Engineer", today, "Verified in browser"
    ),
]
for row in ai_man_tests:
    ws_ai_man.append(row)

# Sheet: Versioning
ws_ver_man = wb_man["Versioning"]
ver_man_tests = [
    (
        "TC-VER-001", "Versioning", "Save as New Version", "Create v2 snapshot from edited draft",
        "Draft v1 edited in Stage 4", "Click 'Save as New Version (v2)'",
        "Modified summary & role title", "New version v2 is created; version dropdown shows v1 and v2",
        "v2 created successfully; switching between v1 and v2 in dropdown preserves respective content", "Passed", "", "Manual / UI", "High",
        "Lead Engineer", today, "Verified in browser"
    ),
    (
        "TC-VER-002", "Versioning", "Lock as Final", "Mark approved resume version as Final",
        "Version v2 reviewed and ready", "Click 'Lock as Final'",
        "Version v2", "Status badge transitions to green FINAL pill; in-place editing locked",
        "Status updated to FINAL; subsequent edits automatically create new draft version v3", "Passed", "", "Manual / UI", "High",
        "Lead Engineer", today, "Verified in browser"
    ),
]
for row in ver_man_tests:
    ws_ver_man.append(row)

# Update Manual Index counts
idx_man = wb_man["00_Index"]
for row in idx_man.iter_rows(min_row=2):
    mod = row[0].value
    if mod == "Tailored_Resume":
        row[2].value = 2
        row[3].value = today
        row[4].value = 2
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (2 manual tests passed)"
    elif mod == "AI_Editing":
        row[2].value = 1
        row[3].value = today
        row[4].value = 1
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (1 manual test passed)"
    elif mod == "Versioning":
        row[2].value = 2
        row[3].value = today
        row[4].value = 2
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (2 manual tests passed)"

wb_man.save(man_path)
print("Updated Manual Test Cases Workbook.")
