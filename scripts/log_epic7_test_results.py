import openpyxl
from datetime import datetime

today = datetime.now().strftime("%Y-%m-%d")

# 1. Update Automation Test Cases Workbook
auto_path = "docs/Verity_Automation_Test_Cases.xlsx"
wb_auto = openpyxl.load_workbook(auto_path)

ws_pdf_auto = wb_auto["PDF_Generation"]
pdf_auto_tests = [
    (
        "AUT-PDF-001", "TC-PDF-001", "PDF_Generation", "LaTeX Escaping", "Sanitize all reserved characters to prevent macro injection and syntax breakage",
        "Unit", "Pytest", "Strings containing &, %, $, #, _, {, }, ~, ^, \\",
        "latex_escape(raw_text)", "All reserved characters properly escaped with TeX control sequences",
        "assert r'\\&' in escaped and r'\\$' in escaped and r'\\%' in escaped",
        "backend/tests/test_pdf_generation.py::test_latex_special_character_escaping",
        "Passed", "", "Local / Epic-7", today, "Verified locally"
    ),
    (
        "AUT-PDF-002", "TC-PDF-002", "PDF_Generation", "Compilation & Persistence", "Invoke pdflatex safely and persist PDF & TeX artifacts",
        "Integration / API", "Pytest / TestClient", "Candidate tailored resume draft in DB",
        "POST /api/v1/tailored-resumes/{id}/generate-pdf", "Returns 200 with pdf_artifact metadata and positive file size",
        "assert data['status'] == 'success' and data['pdf_artifact']['file_size'] > 0",
        "backend/tests/test_pdf_generation.py::test_pdf_compilation_and_artifact_retrieval",
        "Passed", "", "Local / Epic-7", today, "Verified locally"
    ),
    (
        "AUT-PDF-003", "TC-PDF-003", "PDF_Generation", "Document Streaming", "Stream binary PDF and TeX source for preview and download",
        "API", "Pytest / TestClient", "Compiled resume PDF and TeX artifacts",
        "GET /api/v1/tailored-resumes/{id}/pdf and /tex", "Returns application/pdf and text/plain with proper file content",
        "assert stream_res.headers['content-type'] == 'application/pdf'",
        "backend/tests/test_pdf_generation.py::test_pdf_compilation_and_artifact_retrieval",
        "Passed", "", "Local / Epic-7", today, "Verified locally"
    ),
    (
        "AUT-PDF-004", "TC-PDF-004", "PDF_Generation", "Security Isolation", "Prevent unauthorized user from accessing or compiling PDF",
        "Security", "Pytest / TestClient", "User 1 compiled resume artifact",
        "User 2 attempts GET on User 1 PDF and artifacts", "Returns 404 Not Found",
        "assert client.get('/tailored-resumes/{id}/pdf', headers=u2_headers).status_code == 404",
        "backend/tests/test_pdf_generation.py::test_pdf_generation_user_isolation",
        "Passed", "", "Local / Epic-7", today, "Verified locally"
    ),
]

for row in pdf_auto_tests:
    ws_pdf_auto.append(row)

idx_auto = wb_auto["00_Index"]
for row in idx_auto.iter_rows(min_row=2):
    mod = row[0].value
    if mod == "PDF_Generation":
        row[2].value = 4
        row[3].value = today
        row[4].value = 4
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (4 automated tests passed)"

wb_auto.save(auto_path)
print("Updated Automation Test Cases Workbook for Epic 7.")

# 2. Update Manual Test Cases Workbook
man_path = "docs/Verity_Manual_Test_Cases.xlsx"
wb_man = openpyxl.load_workbook(man_path)

ws_pdf_man = wb_man["PDF_Generation"]
pdf_man_tests = [
    (
        "TC-PDF-001", "PDF_Generation", "ATS Resume Compilation", "Compile tailored resume into PDF from Workspace Stage 5",
        "Tailored resume active in Stage 4/5", "Click 'Compile ATS-Ready PDF'",
        "Active tailored resume v1", "Document compiles via local MiKTeX pdflatex; inline iframe renders PDF preview",
        "PDF compiled cleanly and displayed inline in Stage 5 preview iframe", "Passed", "", "Manual / UI", "High",
        "Lead Engineer", today, "Verified in browser"
    ),
    (
        "TC-PDF-002", "PDF_Generation", "Download Artifacts", "Download compiled PDF and LaTeX source files",
        "Compiled PDF preview available in Stage 5", "Click 'Download PDF' and 'LaTeX Source' buttons",
        "Resume_Stripe Platforms.pdf and resume_source.tex", "Browser initiates direct download with correct filenames and binary content",
        "Files downloaded and verified locally in PDF viewer", "Passed", "", "Manual / UI", "High",
        "Lead Engineer", today, "Verified in browser"
    ),
]

for row in pdf_man_tests:
    ws_pdf_man.append(row)

idx_man = wb_man["00_Index"]
for row in idx_man.iter_rows(min_row=2):
    mod = row[0].value
    if mod == "PDF_Generation":
        row[2].value = 2
        row[3].value = today
        row[4].value = 2
        row[5].value = 0
        row[6].value = 0
        row[7].value = "Completed (2 manual tests passed)"

wb_man.save(man_path)
print("Updated Manual Test Cases Workbook for Epic 7.")
