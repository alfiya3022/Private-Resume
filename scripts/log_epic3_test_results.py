import openpyxl
from datetime import datetime

today = datetime.now().strftime("%Y-%m-%d")

# 1. Update Automation Test Cases
auto_path = "docs/PrivateResume_Automation_Test_Cases.xlsx"
wb_auto = openpyxl.load_workbook(auto_path)
ws_job_auto = wb_auto["Workspace_Job_Input"]

job_tests = [
    (
        "AUT-JOB-001", "TC-JOB-001", "Workspace_Job_Input", "Job Creation", "Create and persist target job record",
        "API", "Pytest / TestClient", "Authenticated candidate", "JobCreate payload with TechNova Corp",
        "POST /api/v1/jobs", "Returns 201 Created and persists job record with active status",
        "assert response.status_code == 201 and data['company'] == 'TechNova Corp'",
        "backend/tests/test_jobs.py::test_create_job_success",
        "Passed", "", "Local / Epic-3", today, "Verified locally"
    ),
    (
        "AUT-JOB-002", "TC-JOB-002", "Workspace_Job_Input", "Validation", "Enforce non-empty company, title, and JD text",
        "API / Negative", "Pytest / TestClient", "Authenticated candidate", "Payloads with blank title and blank JD",
        "POST /api/v1/jobs with invalid/empty payloads", "Returns 422 Unprocessable Entity",
        "assert r1.status_code == 422 and r2.status_code == 422",
        "backend/tests/test_jobs.py::test_create_job_empty_validation",
        "Passed", "", "Local / Epic-3", today, "Verified locally"
    ),
    (
        "AUT-JOB-003", "TC-JOB-003", "Workspace_Job_Input", "Job Retrieval", "List jobs and retrieve job by ID",
        "API", "Pytest / TestClient", "Jobs persisted in DB", "Job IDs",
        "GET /api/v1/jobs and GET /api/v1/jobs/{id}", "Returns 200 OK with correct jobs list and job detail",
        "assert list_res.status_code == 200 and get_res.json()['company'] == 'Company A'",
        "backend/tests/test_jobs.py::test_list_and_get_jobs",
        "Passed", "", "Local / Epic-3", today, "Verified locally"
    ),
    (
        "AUT-JOB-004", "TC-JOB-004", "Workspace_Job_Input", "Job Mutation", "Update job fields and delete job record",
        "API", "Pytest / TestClient", "Existing job in DB", "JobUpdate payload",
        "PUT /api/v1/jobs/{id} and DELETE /api/v1/jobs/{id}", "Returns 200 on update, 204 on delete, 404 on subsequent get",
        "assert up_res.status_code == 200 and del_res.status_code == 204",
        "backend/tests/test_jobs.py::test_update_and_delete_job",
        "Passed", "", "Local / Epic-3", today, "Verified locally"
    ),
    (
        "AUT-JOB-005", "TC-JOB-005", "Workspace_Job_Input", "User Isolation", "Cross-user job security and isolation",
        "Security", "Pytest / TestClient", "Two distinct candidates with saved jobs", "User 1 and User 2 tokens",
        "User 2 attempts GET, PUT, DELETE on User 1 job", "Rejected with 404 Not Found; User 2 list is empty",
        "assert 404 on all unauthorized access attempts",
        "backend/tests/test_jobs.py::test_jobs_user_isolation",
        "Passed", "", "Local / Epic-3", today, "Verified locally"
    ),
]

for row in job_tests:
    ws_job_auto.append(row)

# Update Index in Auto workbook
idx_ws_auto = wb_auto["00_Index"]
for row in idx_ws_auto.iter_rows(min_row=2):
    if row[0].value == "Workspace_Job_Input":
        row[2].value = len(job_tests)
        row[3].value = today
        row[4].value = len(job_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "5 automated tests passed (API, Validation, Security)"
        break

wb_auto.save(auto_path)

# 2. Update Manual Test Cases
man_path = "docs/PrivateResume_Manual_Test_Cases.xlsx"
wb_man = openpyxl.load_workbook(man_path)
ws_job_man = wb_man["Workspace_Job_Input"]

job_man_tests = [
    (
        "TC-JOB-001", "Workspace_Job_Input", "Job Intake", "Submit target job description in Workspace",
        "Candidate logged in with active master resume", "Company: Stripe, Title: Backend Engineer, Raw JD text",
        "1. Enter company & role. 2. Paste JD. 3. Click Save & Analyze.",
        "Job record created in database and UI advances to Stage 2 (Analyze JD)",
        "High", "P1", "Functional", "Local Dev",
        "Job persisted, banner indicates active workflow, navigated to Stage 2", "Passed", "", "QA Engineer", today, "", "Verified in browser"
    ),
    (
        "TC-JOB-002", "Workspace_Job_Input", "Validation", "Submit blank job form",
        "Candidate logged in", "Blank inputs",
        "1. Click Save & Analyze without entering title or JD.",
        "Form submission blocked with validation message",
        "Medium", "P2", "Validation", "Local Dev",
        "Client and backend validation trigger, preventing invalid persistence", "Passed", "", "QA Engineer", today, "", "Verified in browser"
    ),
]

for row in job_man_tests:
    ws_job_man.append(row)

# Update Index in Manual workbook
idx_ws_man = wb_man["00_Index"]
for row in idx_ws_man.iter_rows(min_row=2):
    if row[0].value == "Workspace_Job_Input":
        row[2].value = len(job_man_tests)
        row[3].value = today
        row[4].value = len(job_man_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "2 manual test cases executed and passed"
        break

wb_man.save(man_path)
print("Updated Epic 3 test cases successfully!")
