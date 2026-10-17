import openpyxl
from datetime import datetime

today = datetime.now().strftime("%Y-%m-%d")

# 1. Update Automation Test Cases
auto_path = "docs/PrivateResume_Automation_Test_Cases.xlsx"
wb_auto = openpyxl.load_workbook(auto_path)

# Match_Analysis sheet
ws_match_auto = wb_auto["Match_Analysis"]
match_tests = [
    (
        "AUT-MATCH-001", "TC-MATCH-001", "Match_Analysis", "Deterministic Alignment", "4-category classification & Estimated Job Match scoring",
        "API", "Pytest / TestClient", "Profile with Python, FastAPI, and Project", "Job requiring Python, FastAPI, and Kubernetes",
        "POST /api/v1/jobs/{id}/match", "Classifies into Matched, Weak, Missing, and Unsupported; computes 0-100% score",
        "assert 'Python' in data['matched_skills'] and 'Kubernetes' in data['missing_skills']",
        "backend/tests/test_matching.py::test_match_job_and_retrieve",
        "Passed", "", "Local / Epic-5", today, "Verified locally"
    ),
    (
        "AUT-MATCH-002", "TC-MATCH-002", "Match_Analysis", "Security / Isolation", "Reject cross-user match access",
        "Security", "Pytest / TestClient", "Two distinct candidates with saved jobs", "User 1 and User 2 tokens",
        "User 2 attempts GET on User 1 match", "Returns 404 Not Found",
        "assert client.get('/jobs/{id}/match', headers=h2).status_code == 404",
        "backend/tests/test_matching.py::test_match_unauthorized_isolation",
        "Passed", "", "Local / Epic-5", today, "Verified locally"
    ),
]

for row in match_tests:
    ws_match_auto.append(row)

# Evidence sheet
ws_evi_auto = wb_auto["Evidence"]
evi_tests = [
    (
        "AUT-EVI-001", "TC-EVI-001", "Evidence", "Evidence Drawer", "Link candidate projects & experience to matched skills",
        "Unit / API", "Pytest / TestClient", "Profile with projects & experiences", "POST /api/v1/jobs/{id}/match",
        "Inspect evidence_map in match response", "Every matched skill contains citations to exact project titles",
        "assert any('API Backend Engine' in src for src in data['evidence_map']['Python'])",
        "backend/tests/test_matching.py::test_match_job_and_retrieve",
        "Passed", "", "Local / Epic-5", today, "Verified locally"
    ),
    (
        "AUT-EVI-002", "TC-EVI-002", "Evidence", "Integrity Check", "Detect unsupported claims without project/work backing",
        "Unit / API", "Pytest / TestClient", "Skill claimed as Expert with 0 projects", "POST /api/v1/jobs/{id}/match",
        "Inspect unsupported_claims in response", "Identifies unbacked claims with warning instead of silent approval",
        "assert any('GraphQL' in claim for claim in data['unsupported_claims'])",
        "backend/tests/test_matching.py::test_match_job_and_retrieve",
        "Passed", "", "Local / Epic-5", today, "Verified locally"
    ),
]

for row in evi_tests:
    ws_evi_auto.append(row)

# Update Index in Auto workbook
idx_ws_auto = wb_auto["00_Index"]
for row in idx_ws_auto.iter_rows(min_row=2):
    if row[0].value == "Match_Analysis":
        row[2].value = len(match_tests)
        row[3].value = today
        row[4].value = len(match_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "2 automated tests passed (API, Security)"
    elif row[0].value == "Evidence":
        row[2].value = len(evi_tests)
        row[3].value = today
        row[4].value = len(evi_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "2 automated tests passed (Unit, Integrity)"

wb_auto.save(auto_path)

# 2. Update Manual Test Cases
man_path = "docs/PrivateResume_Manual_Test_Cases.xlsx"
wb_man = openpyxl.load_workbook(man_path)

ws_match_man = wb_man["Match_Analysis"]
match_man_tests = [
    (
        "TC-MATCH-001", "Match_Analysis", "Match Scoring", "View Estimated Job Match in Stage 3",
        "Job analyzed in Stage 2", "Candidate profile evidence",
        "1. Click 'Proceed to Evidence Match'. 2. Observe score card.",
        "Displays 0-100% Estimated Job Match (explicitly not labeled ATS score)",
        "High", "P1", "Functional", "Local Dev",
        "Score gauge, matched, weak, missing, and unsupported counts rendered correctly", "Passed", "", "QA Engineer", today, "", "Verified in browser"
    ),
]
for row in match_man_tests:
    ws_match_man.append(row)

ws_evi_man = wb_man["Evidence"]
evi_man_tests = [
    (
        "TC-EVI-001", "Evidence", "Evidence Drawer", "Inspect evidence backing matched skills",
        "Stage 3 active with match results", "None",
        "1. View Matched Skills list. 2. Observe evidence citation below skill name.",
        "Skill citations show exact project names (e.g. '↳ Project: API Backend Engine')",
        "High", "P1", "Functional", "Local Dev",
        "Evidence citations verify candidate claims against actual projects", "Passed", "", "QA Engineer", today, "", "Verified in browser"
    ),
    (
        "TC-EVI-002", "Evidence", "Non-Fabrication", "Verify missing skills warning banner",
        "Stage 3 active with missing skills", "None",
        "1. Observe Missing Skills section.",
        "Displays prominent alert stating missing skills are not instructions to invent experience",
        "High", "P1", "Usability / Integrity", "Local Dev",
        "Non-fabrication banner prominently visible with clear notice", "Passed", "", "QA Engineer", today, "", "Verified in browser"
    ),
]
for row in evi_man_tests:
    ws_evi_man.append(row)

# Update Index in Manual workbook
idx_ws_man = wb_man["00_Index"]
for row in idx_ws_man.iter_rows(min_row=2):
    if row[0].value == "Match_Analysis":
        row[2].value = len(match_man_tests)
        row[3].value = today
        row[4].value = len(match_man_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "1 manual test case executed and passed"
    elif row[0].value == "Evidence":
        row[2].value = len(evi_man_tests)
        row[3].value = today
        row[4].value = len(evi_man_tests)
        row[5].value = 0
        row[6].value = 0
        row[7].value = "2 manual test cases executed and passed"

wb_man.save(man_path)
print("Updated Epic 5 test cases successfully!")
